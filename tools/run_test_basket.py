#!/usr/bin/env python3
"""
Run test basket questions against the current agent and save a report.

Usage:
    python tools/run_test_basket.py
    python tools/run_test_basket.py --xlsx test_data/test_questions.xlsx --limit 5
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import statistics
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from src.agent.agent_graph.agent_graph import app


STOPWORDS_RU = {
    "и",
    "в",
    "во",
    "на",
    "по",
    "с",
    "со",
    "что",
    "как",
    "к",
    "из",
    "для",
    "не",
    "это",
    "а",
    "но",
    "же",
    "или",
    "о",
    "об",
    "от",
    "у",
    "за",
    "над",
    "под",
    "при",
    "его",
    "ее",
    "их",
}


@dataclass
class RunRow:
    row_id: int
    question: str
    expected: str
    actual: str
    status: str
    coverage: float
    duration_sec: float
    error: str


def normalize_tokens(text: str) -> set[str]:
    clean = re.sub(r"[^a-zA-Zа-яА-Я0-9\s]", " ", text.lower())
    tokens = [t for t in clean.split() if len(t) >= 4 and t not in STOPWORDS_RU]
    return set(tokens)


def compute_coverage(expected: str, actual: str) -> float:
    exp_tokens = normalize_tokens(expected)
    act_tokens = normalize_tokens(actual)
    if not exp_tokens:
        return 0.0
    return len(exp_tokens.intersection(act_tokens)) / len(exp_tokens)


def classify(actual: str, error: str, coverage: float, expected: str) -> str:
    if error:
        return "runtime_error"
    if not expected.strip():
        return "no_gold"
    lower = actual.lower()
    if ("ошибка" in lower or "не удалось" in lower) and coverage < 0.15:
        return "failed_answer"
    if coverage >= 0.35:
        return "good_or_partial"
    if coverage >= 0.2:
        return "partial"
    return "weak"


def iter_questions(xlsx_path: Path) -> Iterable[tuple[int, str, str]]:
    wb = load_workbook(filename=xlsx_path, data_only=True)
    ws = wb.active
    for idx, (question, expected) in enumerate(
        ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2
    ):
        q = (question or "").strip()
        e = (expected or "").strip()
        if not q:
            continue
        yield idx, q, e


def run_agent(question: str) -> tuple[str, str, float]:
    start = time.perf_counter()
    try:
        state = app.invoke({"user_question": question})
        actual = str(state.get("final_answer", "")).strip()
        err = ""
    except Exception as ex:  # noqa: BLE001
        actual = ""
        err = str(ex)
    elapsed = time.perf_counter() - start
    return actual, err, elapsed


def save_results(rows: list[RunRow], output_dir: Path) -> tuple[Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = output_dir / f"basket_run_{ts}.json"
    csv_path = output_dir / f"basket_run_{ts}.csv"
    summary_path = output_dir / f"basket_run_{ts}_summary.json"

    payload = [r.__dict__ for r in rows]
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "row_id",
                "status",
                "coverage",
                "duration_sec",
                "question",
                "expected",
                "actual",
                "error",
            ],
        )
        writer.writeheader()
        for r in rows:
            writer.writerow(r.__dict__)

    statuses: dict[str, int] = {}
    durations = []
    coverages = []
    for r in rows:
        statuses[r.status] = statuses.get(r.status, 0) + 1
        durations.append(r.duration_sec)
        if r.expected.strip():
            coverages.append(r.coverage)

    summary = {
        "total_rows": len(rows),
        "status_counts": statuses,
        "avg_duration_sec": statistics.mean(durations) if durations else 0,
        "avg_coverage_gold_rows": statistics.mean(coverages) if coverages else 0,
    }
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return json_path, csv_path, summary_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run test basket against current agent")
    parser.add_argument(
        "--xlsx",
        default="test_data/test_questions.xlsx",
        help="Path to xlsx basket file",
    )
    parser.add_argument(
        "--output-dir",
        default="test_data/results",
        help="Directory to store run artifacts",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Optional limit of questions to run (0 = all)",
    )
    parser.add_argument(
        "--only-with-gold",
        action="store_true",
        help="Run only questions that have a non-empty expected answer",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    xlsx_path = Path(args.xlsx)
    output_dir = Path(args.output_dir)

    rows: list[RunRow] = []
    count = 0
    for row_id, question, expected in iter_questions(xlsx_path):
        if args.only_with_gold and not expected:
            continue
        if args.limit and count >= args.limit:
            break
        actual, err, elapsed = run_agent(question)
        coverage = compute_coverage(expected, actual) if expected else 0.0
        status = classify(actual, err, coverage, expected)
        rows.append(
            RunRow(
                row_id=row_id,
                question=question,
                expected=expected,
                actual=actual,
                status=status,
                coverage=round(coverage, 4),
                duration_sec=round(elapsed, 3),
                error=err,
            )
        )
        count += 1
        print(
            f"[{count}] row={row_id} status={status} coverage={coverage:.3f} "
            f"time={elapsed:.2f}s question={question[:80]}"
        )

    json_path, csv_path, summary_path = save_results(rows, output_dir)
    print(f"Saved JSON: {json_path}")
    print(f"Saved CSV: {csv_path}")
    print(f"Saved summary: {summary_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
