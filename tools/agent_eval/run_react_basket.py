from __future__ import annotations

import argparse
import json
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from src.agent.agent_graph.agent_graph import app


@dataclass
class Question:
    question_id: int
    question: str
    reference_answer: str


def load_questions(xlsx_path: Path, limit: int = 0) -> list[Question]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Questions"] if "Questions" in wb.sheetnames else wb.active
    out: list[Question] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        out.append(
            Question(
                question_id=int(row[0]),
                question=str(row[1] or "").strip(),
                reference_answer=str(row[2] or "").strip(),
            )
        )
        if limit and len(out) >= limit:
            break
    return [q for q in out if q.question]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run ReAct agent on basket and save night-run artifacts")
    parser.add_argument(
        "--questions-xlsx",
        default="test_data/test_questions_binary_eval_v3.xlsx",
    )
    parser.add_argument(
        "--output-root",
        default="test_data/results/night_work_runs",
    )
    parser.add_argument("--run-name", default="react_p0_night")
    parser.add_argument("--limit", type=int, default=0)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_root) / f"{args.run_name}_{ts}"
    run_dir.mkdir(parents=True, exist_ok=True)

    questions = load_questions(Path(args.questions_xlsx), limit=args.limit)

    run_config = {
        "run_id": run_dir.name,
        "questions_xlsx": args.questions_xlsx,
        "questions_count": len(questions),
        "started_at": ts,
        "limit": args.limit,
    }
    (run_dir / "run_config.json").write_text(
        json.dumps(run_config, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    answers: dict[int, str] = {}
    results: list[dict[str, Any]] = []
    react_trace_path = run_dir / "react_trace.jsonl"
    tool_calls_path = run_dir / "tool_calls.jsonl"

    with react_trace_path.open("w", encoding="utf-8") as trace_f, tool_calls_path.open(
        "w", encoding="utf-8"
    ) as tools_f:
        for idx, q in enumerate(questions, start=1):
            started = time.perf_counter()
            error = ""
            state: dict[str, Any] = {}
            try:
                state = app.invoke({"user_question": q.question})
            except Exception as exc:  # noqa: BLE001
                error = str(exc)

            elapsed = round((time.perf_counter() - started) * 1000, 2)
            answer = str((state or {}).get("final_answer", "")).strip()
            answers[q.question_id] = answer

            retrieval_trace = (state or {}).get("retrieval_trace", []) or []
            for item in retrieval_trace:
                trace_f.write(
                    json.dumps(
                        {
                            "question_id": q.question_id,
                            "question": q.question,
                            "trace": item,
                        },
                        ensure_ascii=False,
                    )
                    + "\n"
                )

            tool_calls = (state or {}).get("tool_calls", []) or []
            for item in tool_calls:
                tools_f.write(
                    json.dumps(
                        {
                            "question_id": q.question_id,
                            "question": q.question,
                            "tool_call": item,
                        },
                        ensure_ascii=False,
                    )
                    + "\n"
                )

            results.append(
                {
                    "question_id": q.question_id,
                    "question": q.question,
                    "reference_answer": q.reference_answer,
                    "answer": answer,
                    "duration_ms": elapsed,
                    "error": error,
                    "tool_calls_count": len(tool_calls),
                    "retrieval_trace_count": len(retrieval_trace),
                }
            )
            print(
                f"[{idx}/{len(questions)}] qid={q.question_id} "
                f"tool_calls={len(tool_calls)} trace={len(retrieval_trace)} error={'yes' if error else 'no'}"
            )

    (run_dir / "answers.json").write_text(
        json.dumps(answers, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (run_dir / "basket_results.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (run_dir / "issues_and_actions.md").write_text(
        "# Issues and Actions\n\n- baseline run initialized\n",
        encoding="utf-8",
    )
    (run_dir / "iteration_log.md").write_text(
        "# Iteration Log\n\n",
        encoding="utf-8",
    )

    print(str(run_dir))


if __name__ == "__main__":
    main()

