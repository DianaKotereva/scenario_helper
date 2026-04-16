from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from pydantic import BaseModel, Field, ValidationError

from tools.preprocess_book.utils.llm_factory import create_llm
from tools.agent_eval.prompt import SYSTEM_PROMPT, make_user_prompt


class CheckEval(BaseModel):
    check_id: str
    passed: int = Field(ge=0, le=1)
    reason: str = ""


class EvalResponse(BaseModel):
    checks: list[CheckEval]
    overall_comment: str = ""


@dataclass
class QuestionRow:
    question_id: int
    question: str
    reference_answer: str


@dataclass
class CheckRow:
    question_id: int
    check_id: str
    binary_check: str
    expected: int
    weight: float


def _load_xlsx(path: Path) -> tuple[list[QuestionRow], dict[int, list[CheckRow]]]:
    wb = openpyxl.load_workbook(path)
    qws = wb["Questions"]
    cws = wb["BinaryChecks"]

    questions: list[QuestionRow] = []
    for row in qws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        questions.append(
            QuestionRow(
                question_id=int(row[0]),
                question=str(row[1] or ""),
                reference_answer=str(row[2] or ""),
            )
        )

    checks_by_q: dict[int, list[CheckRow]] = {}
    for row in cws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        item = CheckRow(
            question_id=int(row[0]),
            check_id=str(row[1]),
            binary_check=str(row[2] or ""),
            expected=int(row[3] if row[3] is not None else 1),
            weight=float(row[4] if row[4] is not None else 1.0),
        )
        checks_by_q.setdefault(item.question_id, []).append(item)

    return questions, checks_by_q


def _load_answers_json(path: Path) -> dict[int, str]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    out: dict[int, str] = {}
    if isinstance(payload, dict):
        for k, v in payload.items():
            out[int(k)] = str(v or "")
        return out
    if isinstance(payload, list):
        for item in payload:
            if not isinstance(item, dict):
                continue
            qid = item.get("question_id")
            ans = item.get("answer", item.get("candidate_answer", ""))
            if qid is None:
                continue
            out[int(qid)] = str(ans or "")
    return out


def _safe_parse_json(text: str) -> dict[str, Any]:
    try:
        return json.loads(text)
    except Exception:
        pass
    try:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            return json.loads(text[start : end + 1])
    except Exception:
        pass
    return {}


def _normalize(s: str) -> str:
    return " ".join((s or "").split()).strip().lower()


def evaluate_question_with_llm(
    llm: Any,
    question: QuestionRow,
    checks: list[CheckRow],
    candidate_answer: str,
    retries: int,
) -> EvalResponse:
    checks_payload = [
        {
            "check_id": c.check_id,
            "binary_check": c.binary_check,
            "expected": c.expected,
        }
        for c in checks
    ]

    last_error = ""
    for _ in range(max(1, retries)):
        user_prompt = make_user_prompt(
            question_id=question.question_id,
            question=question.question,
            reference_answer=question.reference_answer,
            candidate_answer=candidate_answer,
            checks_payload=checks_payload,
        )
        resp = llm.invoke(
            [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": user_prompt},
            ]
        )
        content = resp.content if hasattr(resp, "content") else str(resp)
        parsed = _safe_parse_json(content)
        try:
            return EvalResponse.model_validate(parsed)
        except ValidationError as e:
            last_error = str(e)

    return EvalResponse(
        checks=[CheckEval(check_id=c.check_id, passed=0, reason=f"parse_error: {last_error[:180]}") for c in checks],
        overall_comment="llm_parse_failed",
    )


def _evaluate(
    *,
    llm: Any,
    questions: list[QuestionRow],
    checks_by_q: dict[int, list[CheckRow]],
    answers_by_q: dict[int, str],
    retries: int,
    force_llm: bool,
) -> dict[str, Any]:
    per_question = []
    total_expected = 0.0
    total_passed = 0.0

    for q in questions:
        checks = checks_by_q.get(q.question_id, [])
        candidate = answers_by_q.get(q.question_id, q.reference_answer)
        exact = _normalize(candidate) == _normalize(q.reference_answer)

        got: dict[str, CheckEval] = {}
        llm_comment = "skipped_llm"

        if force_llm:
            llm_eval = evaluate_question_with_llm(
                llm=llm,
                question=q,
                checks=checks,
                candidate_answer=candidate,
                retries=retries,
            )
            got = {c.check_id: c for c in llm_eval.checks}
            llm_comment = llm_eval.overall_comment

        check_rows = []
        expected_sum = 0.0
        passed_sum = 0.0

        for c in checks:
            expected_sum += c.weight
            total_expected += c.weight

            # hard safety: etalon answer must be 100%
            if exact:
                passed = c.expected
                reason = "exact_match_with_reference"
            else:
                item = got.get(c.check_id)
                passed = int(item.passed) if item is not None else 0
                reason = item.reason if item is not None else "missing_check_from_llm"

            score = c.weight if (passed == c.expected) else 0.0
            passed_sum += score
            total_passed += score

            check_rows.append(
                {
                    "check_id": c.check_id,
                    "expected": c.expected,
                    "passed": passed,
                    "weight": c.weight,
                    "score": score,
                    "binary_check": c.binary_check,
                    "reason": reason,
                }
            )

        per_question.append(
            {
                "question_id": q.question_id,
                "question": q.question,
                "exact_reference_match": exact,
                "score": passed_sum,
                "max_score": expected_sum,
                "percent": round((passed_sum / expected_sum * 100.0), 2) if expected_sum else 0.0,
                "overall_comment": llm_comment,
                "checks": check_rows,
            }
        )

    overall = round((total_passed / total_expected * 100.0), 2) if total_expected else 0.0
    return {
        "overall_percent": overall,
        "overall_score": total_passed,
        "overall_max_score": total_expected,
        "questions": per_question,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Binary evaluator for ReAct agent answers")
    parser.add_argument(
        "--questions-xlsx",
        default="scenatio_helper/test_data/test_questions_binary_eval_v3.xlsx",
    )
    parser.add_argument("--output-dir", default="scenatio_helper/test_data/results")
    parser.add_argument("--run-name", default="agent_eval_reference_deepseek")
    parser.add_argument("--model", default="deepseek-chat")
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument("--force-llm", action="store_true", help="call LLM before hard etalon scoring")
    parser.add_argument(
        "--answers-json",
        default="",
        help="Optional path to candidate answers: dict {question_id: answer} or list of {question_id, answer}",
    )
    args = parser.parse_args()

    questions_xlsx = Path(args.questions_xlsx)
    output_root = Path(args.output_dir)
    run_id = f"{args.run_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    run_dir = output_root / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    questions, checks_by_q = _load_xlsx(questions_xlsx)
    answers_by_q = _load_answers_json(Path(args.answers_json)) if args.answers_json else {}

    llm = create_llm(
        llm_type="deepseek",
        model=args.model,
        temperature=0,
        max_retries=2,
    )

    report = _evaluate(
        llm=llm,
        questions=questions,
        checks_by_q=checks_by_q,
        answers_by_q=answers_by_q,
        retries=args.retries,
        force_llm=args.force_llm,
    )

    (run_dir / "evaluation_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    summary = {
        "run_id": run_id,
        "questions_xlsx": str(questions_xlsx),
        "model": args.model,
        "force_llm": args.force_llm,
        "answers_json": args.answers_json,
        "overall_percent": report["overall_percent"],
        "overall_score": report["overall_score"],
        "overall_max_score": report["overall_max_score"],
    }
    (run_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
