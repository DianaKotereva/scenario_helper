#!/usr/bin/env python3
from __future__ import annotations

import argparse
import contextvars
import csv
import json
import re
import statistics
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from src.llm_core.llm_prompt_base import LLMBase


STOPWORDS_RU = {
    "Рё",
    "в",
    "во",
    "на",
    "по",
    "с",
    "со",
    "что",
    "как",
    "к",
    "из",
    "для",
    "не",
    "это",
    "а",
    "но",
    "же",
    "или",
    "о",
    "об",
    "от",
    "у",
    "за",
    "над",
    "под",
    "при",
    "его",
    "ее",
    "их",
}


@dataclass
class RunRow:
    row_id: int
    question: str
    expected: str
    actual: str
    status: str
    coverage: float
    duration_sec: float
    error: str


CURRENT_CASE_DIR: contextvars.ContextVar[Path | None] = contextvars.ContextVar(
    "CURRENT_CASE_DIR", default=None
)
_CALL_COUNTERS: dict[str, int] = {}
_LOCK = threading.Lock()


def _safe_json(value: Any) -> Any:
    try:
        json.dumps(value, ensure_ascii=False)
        return value
    except Exception:
        return str(value)


def _next_call_id(case_dir: Path) -> int:
    key = str(case_dir)
    with _LOCK:
        _CALL_COUNTERS[key] = _CALL_COUNTERS.get(key, 0) + 1
        return _CALL_COUNTERS[key]


def _write_llm_call(
    *,
    mode: str,
    self_obj: Any,
    chain_input: Any,
    output: Any = None,
    error: str | None = None,
    elapsed_sec: float = 0.0,
) -> None:
    case_dir = CURRENT_CASE_DIR.get()
    if case_dir is None:
        return
    case_dir.mkdir(parents=True, exist_ok=True)
    call_id = _next_call_id(case_dir)
    payload = {
        "call_id": call_id,
        "mode": mode,
        "class": self_obj.__class__.__name__,
        "elapsed_sec": round(elapsed_sec, 4),
        "input": _safe_json(chain_input),
        "output": _safe_json(output),
        "error": error or "",
    }
    out = case_dir / f"{call_id:04d}_{mode}_{self_obj.__class__.__name__}.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def patch_llm_logging() -> None:
    original_invoke = LLMBase.invoke
    original_ainvoke = LLMBase.ainvoke
    original_process_one = LLMBase._process_one

    def invoke_with_log(self: LLMBase, **kwargs: dict[str, Any]) -> Any:
        chain_input = self.make_user_prompt(**kwargs)
        start = time.perf_counter()
        try:
            result = self.make_llm_chain().invoke(chain_input)
            processed = self._process_output(result)
            _write_llm_call(
                mode="sync_invoke",
                self_obj=self,
                chain_input=chain_input,
                output=processed,
                elapsed_sec=time.perf_counter() - start,
            )
            return processed
        except Exception as ex:  # noqa: BLE001
            _write_llm_call(
                mode="sync_invoke",
                self_obj=self,
                chain_input=chain_input,
                error=str(ex),
                elapsed_sec=time.perf_counter() - start,
            )
            raise

    async def ainvoke_with_log(self: LLMBase, **kwargs: dict[str, Any]) -> Any:
        chain_input = self.make_user_prompt(**kwargs)
        start = time.perf_counter()
        try:
            result = await self.make_llm_chain().ainvoke(chain_input)
            processed = self._process_output(result)
            _write_llm_call(
                mode="async_invoke",
                self_obj=self,
                chain_input=chain_input,
                output=processed,
                elapsed_sec=time.perf_counter() - start,
            )
            return processed
        except Exception as ex:  # noqa: BLE001
            _write_llm_call(
                mode="async_invoke",
                self_obj=self,
                chain_input=chain_input,
                error=str(ex),
                elapsed_sec=time.perf_counter() - start,
            )
            raise

    async def process_one_with_log(
        self: LLMBase, input_data: dict, semaphore: Any
    ) -> Any:
        start = time.perf_counter()
        try:
            res = await original_process_one(self, input_data, semaphore)
            _write_llm_call(
                mode="batch_process_one",
                self_obj=self,
                chain_input=input_data,
                output=res,
                elapsed_sec=time.perf_counter() - start,
            )
            return res
        except Exception as ex:  # noqa: BLE001
            _write_llm_call(
                mode="batch_process_one",
                self_obj=self,
                chain_input=input_data,
                error=str(ex),
                elapsed_sec=time.perf_counter() - start,
            )
            raise

    # Keep references alive for debugging; behavior intentionally overrides originals.
    LLMBase.invoke = invoke_with_log  # type: ignore[assignment]
    LLMBase.ainvoke = ainvoke_with_log  # type: ignore[assignment]
    LLMBase._process_one = process_one_with_log  # type: ignore[assignment]
    _ = (original_invoke, original_ainvoke)


def normalize_tokens(text: str) -> set[str]:
    clean = re.sub(r"[^a-zA-Zа-яА-Я0-9\s]", " ", text.lower())
    tokens = [t for t in clean.split() if len(t) >= 4 and t not in STOPWORDS_RU]
    return set(tokens)


def compute_coverage(expected: str, actual: str) -> float:
    exp_tokens = normalize_tokens(expected)
    act_tokens = normalize_tokens(actual)
    if not exp_tokens:
        return 0.0
    return len(exp_tokens.intersection(act_tokens)) / len(exp_tokens)


def classify(actual: str, error: str, coverage: float, expected: str) -> str:
    if error:
        return "runtime_error"
    if not expected.strip():
        return "no_gold"
    lower = actual.lower()
    if ("ошибка" in lower or "не удалось" in lower) and coverage < 0.15:
        return "failed_answer"
    if coverage >= 0.35:
        return "good_or_partial"
    if coverage >= 0.2:
        return "partial"
    return "weak"


def iter_questions(xlsx_path: Path) -> Iterable[tuple[int, str, str]]:
    wb = load_workbook(filename=xlsx_path, data_only=True)
    ws = wb.active
    for idx, (question, expected) in enumerate(
        ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2
    ):
        q = (question or "").strip()
        e = (expected or "").strip()
        if not q:
            continue
        yield idx, q, e


def save_results(rows: list[RunRow], output_dir: Path) -> tuple[Path, Path, Path, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = output_dir / f"basket_run_{ts}.json"
    csv_path = output_dir / f"basket_run_{ts}.csv"
    summary_path = output_dir / f"basket_run_{ts}_summary.json"

    payload = [r.__dict__ for r in rows]
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "row_id",
                "status",
                "coverage",
                "duration_sec",
                "question",
                "expected",
                "actual",
                "error",
            ],
        )
        writer.writeheader()
        for r in rows:
            writer.writerow(r.__dict__)

    statuses: dict[str, int] = {}
    durations = []
    coverages = []
    for r in rows:
        statuses[r.status] = statuses.get(r.status, 0) + 1
        durations.append(r.duration_sec)
        if r.expected.strip():
            coverages.append(r.coverage)

    summary = {
        "total_rows": len(rows),
        "status_counts": statuses,
        "avg_duration_sec": statistics.mean(durations) if durations else 0,
        "avg_coverage_gold_rows": statistics.mean(coverages) if coverages else 0,
    }
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return json_path, csv_path, summary_path, ts


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run test basket + log every LLM call per question"
    )
    parser.add_argument("--xlsx", default="test_data/test_questions.xlsx")
    parser.add_argument("--output-dir", default="test_data/results")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--only-with-gold", action="store_true")
    parser.add_argument(
        "--llm-log-root",
        default="test_data/results/llm_call_logs",
        help="Root directory for per-run/per-question LLM call logs",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    xlsx_path = Path(args.xlsx)
    output_dir = Path(args.output_dir)
    llm_log_root = Path(args.llm_log_root)

    patch_llm_logging()
    from src.agent.agent_graph.agent_graph import app  # noqa: WPS433

    rows: list[RunRow] = []
    count = 0
    run_ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_log_dir = llm_log_root / f"basket_run_{run_ts}"
    run_log_dir.mkdir(parents=True, exist_ok=True)

    for row_id, question, expected in iter_questions(xlsx_path):
        if args.only_with_gold and not expected:
            continue
        if args.limit and count >= args.limit:
            break

        case_dir = run_log_dir / f"row_{row_id:03d}"
        token = CURRENT_CASE_DIR.set(case_dir)
        start = time.perf_counter()
        try:
            state = app.invoke({"user_question": question})
            actual = str(state.get("final_answer", "")).strip()
            err = ""
        except Exception as ex:  # noqa: BLE001
            actual = ""
            err = str(ex)
        finally:
            CURRENT_CASE_DIR.reset(token)
        elapsed = time.perf_counter() - start

        coverage = compute_coverage(expected, actual) if expected else 0.0
        status = classify(actual, err, coverage, expected)
        rows.append(
            RunRow(
                row_id=row_id,
                question=question,
                expected=expected,
                actual=actual,
                status=status,
                coverage=round(coverage, 4),
                duration_sec=round(elapsed, 3),
                error=err,
            )
        )
        count += 1
        print(
            f"[{count}] row={row_id} status={status} coverage={coverage:.3f} "
            f"time={elapsed:.2f}s logs={case_dir.name}"
        )

    json_path, csv_path, summary_path, ts = save_results(rows, output_dir)
    final_log_dir = llm_log_root / f"basket_run_{ts}"
    if final_log_dir.exists():
        # avoid collision in unlikely same-second case
        final_log_dir = llm_log_root / f"basket_run_{ts}_llm"
    run_log_dir.rename(final_log_dir)

    print(f"Saved JSON: {json_path}")
    print(f"Saved CSV: {csv_path}")
    print(f"Saved summary: {summary_path}")
    print(f"Saved LLM logs: {final_log_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
