#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import time
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Tuple

from openpyxl import load_workbook

from src.agent.agent_graph.agent_graph import app
from tools.prompt_debug.error_analyzer import (
    CaseAnalysis,
    analyze_case,
    save_analysis,
    summarize_analyses,
)
from tools.prompt_debug.instrumentation import AgentDiagnosticsInstrumentation
from tools.prompt_debug.llm_judge import judge_case_with_llm
from tools.prompt_debug.stage_logger import CaseMeta, StageLogger


@dataclass
class CaseRun:
    case_id: str
    row_id: int
    question: str
    expected: str
    actual: str
    runtime_error: str
    duration_sec: float
    final_state: dict


def iter_questions(xlsx_path: Path) -> Iterable[Tuple[int, str, str]]:
    wb = load_workbook(filename=xlsx_path, data_only=True)
    ws = wb.active
    for idx, (question, expected) in enumerate(
        ws.iter_rows(min_row=2, max_col=2, values_only=True),
        start=2,
    ):
        q = (question or "").strip()
        e = (expected or "").strip()
        if q:
            yield idx, q, e


def run_case(question: str) -> Tuple[dict, str, float]:
    started = time.perf_counter()
    try:
        state = app.invoke({"user_question": question})
        runtime_error = ""
    except Exception as ex:  # noqa: BLE001
        state = {}
        runtime_error = str(ex)
    elapsed = time.perf_counter() - started
    return state, runtime_error, elapsed


def save_case_artifact(
    case_run: CaseRun,
    events: List[dict],
    analysis: CaseAnalysis,
    output_dir: Path,
) -> None:
    payload = {
        "case": asdict(case_run),
        "events": events,
        "analysis": analysis.as_dict(),
    }
    file_path = output_dir / f"{case_run.case_id}.json"
    file_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def save_summary_md(
    output_path: Path,
    summary: dict,
    runs: List[CaseRun],
    analyses: List[CaseAnalysis],
) -> None:
    by_case = {a.case_id: a for a in analyses}
    lines = [
        "# Prompt Diagnostics Summary",
        "",
        f"- Total cases: {len(runs)}",
        f"- Avg coverage: {summary.get('avg_coverage', 0)}",
        f"- Worst stage: {summary.get('worst_stage', 'ok')} ({summary.get('worst_stage_count', 0)})",
        "",
        "## Stage Counts",
    ]
    for stage, count in sorted(summary.get("stage_counts", {}).items(), key=lambda x: x[0]):
        lines.append(f"- {stage}: {count}")

    lines.extend(["", "## Per-Case", ""])
    for run in runs:
        a = by_case[run.case_id]
        lines.append(f"### {run.case_id}")
        lines.append(f"- row_id: {run.row_id}")
        lines.append(f"- stage: {a.stage}")
        lines.append(f"- severity: {a.severity}")
        lines.append(f"- confidence: {a.confidence}")
        lines.append(f"- duration_sec: {run.duration_sec:.2f}")
        lines.append(f"- question: {run.question}")
        lines.append(f"- reason: {a.reasons[0] if a.reasons else '-'}")
        lines.append("")

    output_path.write_text("\n".join(lines), encoding="utf-8")


def save_runs_csv(runs: List[CaseRun], analyses: List[CaseAnalysis], output_path: Path) -> None:
    by_case = {a.case_id: a for a in analyses}
    with output_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "case_id",
                "row_id",
                "stage",
                "severity",
                "confidence",
                "coverage",
                "duration_sec",
                "question",
                "expected",
                "actual",
                "runtime_error",
            ],
        )
        writer.writeheader()
        for run in runs:
            analysis = by_case[run.case_id]
            writer.writerow(
                {
                    "case_id": run.case_id,
                    "row_id": run.row_id,
                    "stage": analysis.stage,
                    "severity": analysis.severity,
                    "confidence": analysis.confidence,
                    "coverage": analysis.metrics.get("coverage", 0.0),
                    "duration_sec": round(run.duration_sec, 3),
                    "question": run.question,
                    "expected": run.expected,
                    "actual": run.actual,
                    "runtime_error": run.runtime_error,
                }
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run full stage diagnostics for the test basket.")
    parser.add_argument("--xlsx", default="test_data/test_questions.xlsx", help="Path to basket XLSX.")
    parser.add_argument("--output-dir", default="test_data/results", help="Root output dir.")
    parser.add_argument("--limit", type=int, default=0, help="Optional limit (0 = all).")
    parser.add_argument(
        "--only-with-gold",
        action="store_true",
        help="Run only rows that have non-empty expected answer.",
    )
    parser.add_argument(
        "--no-llm-judge",
        action="store_true",
        help="Disable LLM-as-a-judge and use only heuristic analyzer.",
    )
    parser.add_argument(
        "--row-ids",
        default="",
        help="Comma-separated Excel row ids to run (example: 2,17,19).",
    )
    parser.add_argument(
        "--save-trace",
        action="store_true",
        help="Compatibility flag: traces are always saved into events.jsonl/cases/*.json.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_id = f"prompt_debug_{ts}"
    output_root = Path(args.output_dir) / run_id
    cases_dir = output_root / "cases"
    cases_dir.mkdir(parents=True, exist_ok=True)

    logger = StageLogger(output_dir=output_root, run_id=run_id)
    runs: List[CaseRun] = []
    analyses: List[CaseAnalysis] = []

    with AgentDiagnosticsInstrumentation(logger) as instrumentation:
        processed = 0
        selected_row_ids = set()
        if args.row_ids.strip():
            selected_row_ids = {
                int(x.strip())
                for x in args.row_ids.split(",")
                if x.strip().isdigit()
            }
        for row_id, question, expected in iter_questions(Path(args.xlsx)):
            if selected_row_ids and row_id not in selected_row_ids:
                continue
            if args.only_with_gold and not expected:
                continue
            if args.limit and processed >= args.limit:
                break

            case_id = f"row_{row_id}"
            logger.register_case(
                CaseMeta(case_id=case_id, row_id=row_id, question=question, expected=expected)
            )
            instrumentation.set_case(case_id)

            final_state, runtime_error, duration_sec = run_case(question)
            actual = str(final_state.get("final_answer", "")).strip() if isinstance(final_state, dict) else ""

            case_run = CaseRun(
                case_id=case_id,
                row_id=row_id,
                question=question,
                expected=expected,
                actual=actual,
                runtime_error=runtime_error,
                duration_sec=duration_sec,
                final_state=final_state if isinstance(final_state, dict) else {},
            )
            runs.append(case_run)

            case_events = logger.events_by_case.get(case_id, [])
            analysis = analyze_case(
                case_id=case_id,
                expected=expected,
                actual=actual,
                events=case_events,
                runtime_error=runtime_error,
            )

            if not args.no_llm_judge:
                try:
                    judge = judge_case_with_llm(
                        question=question,
                        expected=expected,
                        actual=actual,
                        events=case_events,
                    )
                    analysis.metrics["llm_judge"] = judge.as_dict()
                    # Trust the judge when confidence is high.
                    if judge.confidence >= 0.7:
                        analysis.stage = judge.stage
                        analysis.confidence = judge.confidence
                        analysis.reasons.insert(0, f"LLM judge: {judge.reason}")
                        analysis.metrics["retrieval_relevance_llm"] = (
                            judge.retrieval_relevance
                        )
                        analysis.metrics["graph_relevance_llm"] = judge.graph_relevance
                except Exception as judge_ex:  # noqa: BLE001
                    analysis.metrics["llm_judge_error"] = str(judge_ex)

            analyses.append(analysis)
            save_case_artifact(case_run, case_events, analysis, cases_dir)

            processed += 1
            print(
                f"[{processed}] {case_id} stage={analysis.stage} "
                f"coverage={analysis.metrics.get('coverage', 0):.3f} "
                f"time={duration_sec:.2f}s"
            )

    summary = summarize_analyses(analyses)
    save_analysis(output_root / "analysis.json", analyses, summary)
    save_runs_csv(runs, analyses, output_root / "summary.csv")
    save_summary_md(output_root / "SUMMARY.md", summary, runs, analyses)

    print(f"Saved diagnostics to: {output_root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
